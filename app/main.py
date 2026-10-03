
from fastapi import FastAPI, Request, Form, BackgroundTasks
from fastapi.responses import HTMLResponse, FileResponse
from fastapi.templating import Jinja2Templates
from pathlib import Path
from datetime import date, datetime
from calendar import monthrange
from tempfile import TemporaryDirectory
from copy import deepcopy
import os, re, shutil, subprocess, uuid

from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = BASE_DIR / "template.docx"
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

app = FastAPI(title="Quotation PDF Generator")


def money(v):
    return f"{float(v):.2f}"


def replace_in_paragraph(paragraph, replacements):
    """Replace placeholders even when Word split them across multiple runs."""
    full = "".join(run.text or "" for run in paragraph.runs)
    if not full:
        return
    new_text = full
    for old, new in replacements.items():
        new_text = new_text.replace(old, str(new))
    if new_text == full:
        return

    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def replace_everywhere(doc, replacements):
    for p in doc.paragraphs:
        replace_in_paragraph(p, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, replacements)


def set_cell_text(cell, text):
    # Preserve the existing paragraph/run formatting as much as possible.
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = str(text)
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(str(text))
    for extra in cell.paragraphs[1:]:
        for r in extra.runs:
            r.text = ""


def clone_row(table, source_row):
    new_tr = deepcopy(source_row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def remove_row(table, row):
    table._tbl.remove(row._tr)


def populate_service_table(table, items, prefix, has_denom=False):
    """
    Template tables have two sample item rows. The first item row is used
    as the formatting prototype; all sample rows are removed and cloned.
    """
    if len(table.rows) < 4:
        return

    prototype = deepcopy(table.rows[3]._tr)
    # Remove all existing item rows (rows 3 onward).
    while len(table.rows) > 3:
        remove_row(table, table.rows[3])

    if not items:
        return

    for idx, item in enumerate(items, start=1):
        new_tr = deepcopy(prototype)
        table._tbl.append(new_tr)
        row = table.rows[-1]

        vals = []
        if prefix == "s":
            vals = [
                idx,
                item["description"],
                item["denom"],
                item["qty"],
                "9.0",
                money(item["gst"]),
                "9.0",
                money(item["gst"]),
                money(item["amt"]),
            ]
        else:
            vals = [
                idx,
                item["description"],
                item["qty"],
                money(item["rate"]),
                "9.0",
                money(item["gst"]),
                "9.0",
                money(item["gst"]),
                money(item["amt"]),
            ]

        for c, value in zip(row.cells, vals):
            set_cell_text(c, value)


def convert_to_pdf(docx_path, out_dir):
    """Convert DOCX to PDF using LibreOffice/soffice."""
    candidates = [
        os.environ.get("LIBREOFFICE_PATH"),
        os.environ.get("SOFFICE_PATH"),
        shutil.which("soffice"),
        shutil.which("libreoffice")
    ]
    soffice = next((p for p in candidates if p and os.path.exists(p)), None)
    if not soffice:
        raise RuntimeError(
            "LibreOffice was not found. Install LibreOffice and either add "
            "'soffice' to PATH or set the LIBREOFFICE_PATH environment variable."
        )

    result = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        timeout=120,
    )
    pdf = Path(out_dir) / (Path(docx_path).stem + ".pdf")
    if result.returncode != 0 or not pdf.exists():
        raise RuntimeError(
            "PDF conversion failed.\nSTDOUT: " + result.stdout + "\nSTDERR: " + result.stderr
        )
    return pdf


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={"request": request}
    )


@app.post("/generate")
async def generate(
    request: Request,
    background_tasks: BackgroundTasks,
    quote_date: str = Form(...),
    city: str = Form(...),
    customer_name: str = Form(...),
    customer_address: str = Form(...),
    customer_cno: str = Form(...),
    storage_date: str = Form(...),
    area: str = Form(...),
    tenure: str = Form(...),
    storage_json: str = Form("[]"),
    handling_json: str = Form("[]"),
    addon_json: str = Form("[]"),
):
    import json

    try:
        qd = datetime.strptime(quote_date, "%Y-%m-%d").date()
    except ValueError:
        return HTMLResponse("Invalid quotation date.", status_code=400)

    try:
        storage = json.loads(storage_json or "[]")
        handling = json.loads(handling_json or "[]")
        addon = json.loads(addon_json or "[]")
    except json.JSONDecodeError:
        return HTMLResponse("Invalid item data.", status_code=400)

    def num(x, field):
        try:
            return float(x)
        except (TypeError, ValueError):
            raise ValueError(f"Invalid {field}: {x}")

    # Server-side calculations: never trust browser-calculated amounts.
    storage_out = []
    for item in storage:
        rate = num(item.get("rate"), "storage rate")
        qty = num(item.get("qty"), "storage quantity")
        amt = rate * qty
        gst = amt * 0.09
        storage_out.append({
            "description": str(item.get("description", "")),
            "denom": str(item.get("denom", "")),
            "qty": qty,
            "rate": rate,
            "amt": amt,
            "gst": gst,
        })

    handling_out = []
    for item in handling:
        rate = num(item.get("rate"), "handling rate")
        qty = num(item.get("qty"), "handling quantity")
        # Per the requested rule: amount = rate (quantity does not multiply it).
        amt = rate
        gst = amt * 0.09
        handling_out.append({
            "description": str(item.get("description", "")),
            "qty": qty,
            "rate": rate,
            "amt": amt,
            "gst": gst,
        })

    addon_out = []
    for item in addon:
        rate = num(item.get("rate"), "add-on rate")
        qty = num(item.get("qty"), "add-on quantity")
        # Per the requested rule: amount = rate.
        amt = rate
        gst = amt * 0.09
        addon_out.append({
            "description": str(item.get("description", "")),
            "qty": qty,
            "rate": rate,
            "amt": amt,
            "gst": gst,
        })

    # The uploaded template currently contains {{storage_date}}, {{Produt_cat}}
    # and the customer fields in the bill-to table; retain those placeholders.
    last_day = date(qd.year, qd.month, monthrange(qd.year, qd.month)[1])

    replacements = {
        "{{date}}": qd.strftime("%d/%m/%Y"),
        "{{city}}": city,
        "{{customer.name}}": customer_name,
        "{{customer.address}}": customer_address,
        "{{customer.cno}}": customer_cno,
        "{{storage_date}}": storage_date,
        "{{Area}}": area,
        "{{tenure}}": tenure,
        "{{last_day_of_date_month}}": last_day.strftime("%d/%m/%Y"),
        "{{last_day_of_ date _month}}": last_day.strftime("%d/%m/%Y"),
        "{{Produt_cat}}": "",
    }

    tmp_path = Path(__import__("tempfile").mkdtemp(prefix="quotation_"))
    try:
        working_docx = tmp_path / f"quotation_{uuid.uuid4().hex}.docx"
        shutil.copy(TEMPLATE_PATH, working_docx)
        doc = Document(working_docx)

        # Tables 1, 2, 3 in the supplied template are the three service tables.
        populate_service_table(doc.tables[1], storage_out, "s", has_denom=True)
        populate_service_table(doc.tables[2], handling_out, "h")
        populate_service_table(doc.tables[3], addon_out, "r")

        replace_everywhere(doc, replacements)
        doc.save(working_docx)

        pdf = convert_to_pdf(working_docx, tmp_path)
        final_pdf = tmp_path / f"Quotation_{qd.strftime('%Y%m%d')}_{uuid.uuid4().hex[:8]}.pdf"
        shutil.copy(pdf, final_pdf)

        # Keep the temporary directory alive until the response has been sent.
        background_tasks.add_task(shutil.rmtree, tmp_path, ignore_errors=True)
        return FileResponse(
            path=str(final_pdf),
            media_type="application/pdf",
            filename=final_pdf.name,
            headers={"Content-Disposition": f'attachment; filename="{final_pdf.name}"'},
        )

    except ValueError as exc:
        shutil.rmtree(tmp_path, ignore_errors=True)
        return HTMLResponse(str(exc), status_code=400)
    except Exception as exc:
        shutil.rmtree(tmp_path, ignore_errors=True)
        return HTMLResponse(
            "Unable to generate quotation PDF: " + str(exc),
            status_code=500,
        )
